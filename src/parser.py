"""Stage 1: Resume parser. Text extraction + structured fields."""
import io
import re
from dataclasses import dataclass, field, asdict

from .skills import extract_skills

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s-]?)?(?:\(?\d{3,5}\)?[\s-]?)?\d{3,5}[\s-]?\d{4,5}")
LINK_RE = re.compile(r"(?:https?://)?(?:www\.)?(?:linkedin\.com|github\.com)/[\w\-/.]+", re.I)

SECTION_HEADINGS = {
    "summary": ["summary", "profile", "objective", "about me", "career objective"],
    "skills": ["skills", "technical skills", "core competencies", "technologies"],
    "experience": ["experience", "work experience", "professional experience", "employment", "internships"],
    "education": ["education", "academic background", "qualifications"],
    "projects": ["projects", "personal projects", "academic projects"],
    "certifications": ["certifications", "certificates", "courses", "achievements", "awards"],
}


@dataclass
class ParsedResume:
    name: str = ""
    email: str = ""
    phone: str = ""
    links: list = field(default_factory=list)
    skills: list = field(default_factory=list)
    sections: dict = field(default_factory=dict)
    raw_text: str = ""

    def to_dict(self):
        return asdict(self)


def extract_text(data: bytes, filename: str) -> str:
    """Read text from PDF, DOCX or TXT bytes."""
    if len(data) > 10 * 1024 * 1024:
        raise ValueError("Upload a file smaller than 10 MB.")
    name = filename.lower()
    if not name.endswith((".pdf", ".docx", ".txt")):
        raise ValueError("Supported files: PDF, DOCX and TXT.")
    if name.endswith(".pdf"):
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join((page.extract_text() or "") for page in pdf.pages)
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return data.decode("utf-8", errors="ignore")


def _heading_key(line: str):
    clean = re.sub(r"[^a-z ]", "", line.lower()).strip()
    if not clean or len(clean.split()) > 4:
        return None
    for key, names in SECTION_HEADINGS.items():
        if clean in names:
            return key
    return None


def split_sections(text: str) -> dict:
    sections, current, buf = {}, "header", []
    for line in text.splitlines():
        key = _heading_key(line)
        if key:
            sections[current] = "\n".join(buf).strip()
            current, buf = key, []
        else:
            buf.append(line)
    sections[current] = "\n".join(buf).strip()
    return {k: v for k, v in sections.items() if v}


def _guess_name(text: str) -> str:
    try:  # optional spaCy upgrade
        import spacy
        nlp = spacy.load("en_core_web_sm")
        for ent in nlp(text[:400]).ents:
            if ent.label_ == "PERSON":
                return ent.text.strip()
    except Exception:
        pass
    for line in text.splitlines()[:6]:
        line = line.strip()
        if line and not EMAIL_RE.search(line) and not any(c.isdigit() for c in line) and 1 < len(line.split()) <= 4:
            return line.title() if line.isupper() else line
    return ""


def parse_resume(text: str) -> ParsedResume:
    email = EMAIL_RE.search(text)
    phone = next((m.group().strip() for m in PHONE_RE.finditer(text) if len(re.sub(r"\D", "", m.group())) >= 10), "")
    return ParsedResume(
        name=_guess_name(text),
        email=email.group() if email else "",
        phone=phone,
        links=sorted(set(LINK_RE.findall(text))),
        skills=extract_skills(text),
        sections=split_sections(text),
        raw_text=text,
    )
