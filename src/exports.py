"""Simple single-column document exports. User text is never treated as HTML."""
from io import BytesIO
import re
from xml.sax.saxutils import escape


def lines(text):
    for line in text.splitlines():
        line = line.strip()
        if not line:
            yield 'blank', ''
        elif line.lower().rstrip(':') in {"summary", "skills", "technical skills", "experience", "education", "projects", "certifications", "professional experience"}:
            yield 'heading', line.rstrip(':')
        elif line.startswith('#'):
            yield 'heading', line.lstrip('#').strip()
        elif line.startswith(('- ', '* ', '• ')):
            yield 'bullet', line[2:]
        else:
            yield 'text', re.sub(r'\*\*(.*?)\*\*', r'\1', line)


def docx_bytes(text, title='Resume', template='Classic'):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    # Remove inherited decorative borders from default Word styles.
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.75)
    style = doc.styles['Normal']
    style.font.name = 'Calibri' if template == 'Modern' else 'Arial'
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(5)
    doc.add_heading(title, 0)
    doc.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Heading 2'].font.color.rgb = RGBColor.from_string('155E75' if template == 'Modern' else '000000')
    if text.splitlines() and text.splitlines()[0].strip() == title.strip():
        text = '\n'.join(text.splitlines()[1:])
    for kind, line in lines(text):
        if kind == 'heading':
            doc.add_heading(line, 2)
        elif kind == 'bullet':
            doc.add_paragraph(line, style='List Bullet')
        elif kind == 'text':
            doc.add_paragraph(line)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def pdf_bytes(text, title='Resume', template='Classic'):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.pagesizes import A4
    import reportlab
    from pathlib import Path
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font_dir = Path(reportlab.__file__).parent / 'fonts'
    if 'ResumeVera' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('ResumeVera', str(font_dir / 'Vera.ttf')))
        pdfmetrics.registerFont(TTFont('ResumeVeraBold', str(font_dir / 'VeraBd.ttf')))
    stream = BytesIO()
    styles = getSampleStyleSheet()
    styles['Title'].fontName = 'ResumeVeraBold'
    styles['Heading2'].fontName = 'ResumeVeraBold'
    normal = ParagraphStyle('ResumeBody', parent=styles['BodyText'], fontName='ResumeVera', fontSize=10,
                            leading=14, spaceAfter=5, alignment=TA_LEFT, splitLongWords=True)
    heading = ParagraphStyle('ResumeHeading', parent=styles['Heading2'], textColor=colors.HexColor('#155e75') if template == 'Modern' else colors.black)
    # Portable built-in font supports Latin-1. Fail clearly for unsupported scripts.
    combined = title + text
    substitutions = {'–': '-', '—': '-', '’': "'", '‘': "'", '“': '"', '”': '"', '•': '-', '…': '...'}
    for source, target in substitutions.items():
        combined = combined.replace(source, target)
        title = title.replace(source, target)
        text = text.replace(source, target)
    try:
        combined.encode('latin-1')
    except UnicodeEncodeError as exc:
        raise ValueError('PDF supports Latin text. Use DOCX for Tamil or other scripts.') from exc
    story = [Paragraph(escape(title), styles['Title']), Spacer(1, 10)]
    if text.splitlines() and text.splitlines()[0].strip() == title.strip():
        text = '\n'.join(text.splitlines()[1:])
    for kind, line in lines(text):
        if kind == 'blank':
            story.append(Spacer(1, 5))
        else:
            story.append(Paragraph(escape(('- ' if kind == 'bullet' else '') + line), heading if kind == 'heading' else normal))
    SimpleDocTemplate(stream, pagesize=A4, leftMargin=48, rightMargin=48, topMargin=40, bottomMargin=40).build(story)
    return stream.getvalue()
